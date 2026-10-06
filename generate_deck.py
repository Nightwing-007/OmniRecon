#!/usr/bin/env python3
"""
OmniRecon Presentation Generator
=================================
Automated slide deck generator using python-pptx.
Produces a modern, cyber-themed, high-contrast technical presentation
tailored for academic reviewers and technical recruiters.

Theme: Deep Charcoal/Navy background, Neon Cyan, Emerald Mint, Coral accents,
clean card-based modular layout with minimal text and punchy bullet points.
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# ==========================================
# COLOR PALETTE SPECIFICATION (DARK CYBER)
# ==========================================
BG_COLOR         = RGBColor(10, 14, 23)     # #0A0E17 (Deep Navy Charcoal)
CARD_BG          = RGBColor(19, 27, 46)     # #131B2E (Dark Slate Card)
CARD_BORDER      = RGBColor(35, 50, 78)     # #23324E (Subtle Structural Border)
CARD_BG_ALT      = RGBColor(15, 23, 40)     # #0F1728 (Slightly Darker Card)

CYAN_ACCENT      = RGBColor(0, 229, 255)    # #00E5FF (Electric Cyan - Primary)
EMERALD_ACCENT   = RGBColor(0, 230, 118)    # #00E676 (Mint Emerald - Success/Active)
CORAL_ACCENT     = RGBColor(255, 64, 129)   # #FF4081 (Vibrant Coral/Pink - Vuln/Alert)
AMBER_ACCENT     = RGBColor(255, 179, 0)    # #FFB300 (Warning Amber)
BLUE_ACCENT      = RGBColor(68, 138, 255)   # #448AFF (Tech Blue)

TEXT_WHITE       = RGBColor(255, 255, 255)  # #FFFFFF (Crisp White Headers)
TEXT_MUTED       = RGBColor(186, 201, 224)  # #BAC9E0 (Readable Light Slate Body)
TEXT_DIM         = RGBColor(118, 139, 173)  # #768BAD (Metadata & Captions)
BADGE_BG         = RGBColor(24, 38, 66)     # #182642 (Pill Badge Background)

FONT_HEADING = "Trebuchet MS"
FONT_BODY    = "Calibri"
FONT_CODE    = "Consolas"


def create_base_slide(prs, category_kicker: str, slide_title: str, subtitle_takeaway: str = None):
    """Creates a slide with the dark theme background, kicker tag, title, and accent line."""
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # 1. Dark Background Shape (Full Bleed)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.color.rgb = BG_COLOR

    # 2. Top Kicker & Slide Title Box
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # Category Kicker (e.g., "// 02. ATTACK SURFACE EXPANSION")
    p_kicker = tf.paragraphs[0]
    p_kicker.text = category_kicker.upper()
    p_kicker.font.name = FONT_CODE
    p_kicker.font.size = Pt(10)
    p_kicker.font.bold = True
    p_kicker.font.color.rgb = CYAN_ACCENT
    p_kicker.space_after = Pt(2)

    # Main Title
    p_title = tf.add_paragraph()
    p_title.text = slide_title
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE

    # Optional Subtitle Takeaway
    if subtitle_takeaway:
        p_title.space_after = Pt(2)
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle_takeaway
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = TEXT_MUTED

    # 3. Accent divider line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.55), Inches(11.733), Inches(0.025))
    line.fill.solid()
    line.fill.fore_color.rgb = CARD_BORDER
    line.line.color.rgb = CARD_BORDER

    # 4. Subtle Footer
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
    ftf = footer_box.text_frame
    ftf.word_wrap = True
    ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0
    p_foot = ftf.paragraphs[0]
    p_foot.text = "OmniRecon  |  Modular Reconnaissance & Attack Surface Toolkit  •  Deepakraj S"
    p_foot.font.name = FONT_CODE
    p_foot.font.size = Pt(9)
    p_foot.font.color.rgb = TEXT_DIM

    return slide


def add_card(slide, left, top, width, height, title=None, badge_text=None,
             badge_color=CYAN_ACCENT, bg_color=CARD_BG, border_color=CARD_BORDER):
    """Renders a modern rounded container card with optional header and badge."""
    # Outer card shape
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)

    content_top = top + Inches(0.2)
    content_height = height - Inches(0.4)

    # If title or badge provided, create top header within card
    if title or badge_text:
        header_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), Inches(0.55))
        htf = header_box.text_frame
        htf.word_wrap = True
        htf.margin_left = htf.margin_top = htf.margin_right = htf.margin_bottom = 0

        p = htf.paragraphs[0]
        if badge_text:
            p.text = f"[{badge_text}] "
            p.font.name = FONT_CODE
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = badge_color

        if title:
            # Append or write title
            run = p.add_run() if badge_text else p
            run.text = title
            run.font.name = FONT_HEADING
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = TEXT_WHITE

        # Adjust remaining content box
        content_top = top + Inches(0.7)
        content_height = height - Inches(0.85)

    # Inner content text box
    tb = slide.shapes.add_textbox(left + Inches(0.25), content_top, width - Inches(0.5), content_height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    return tf


def add_bullet_point(text_frame, title_bold: str, description: str, accent_color=CYAN_ACCENT,
                     font_size=12, is_first=False, icon="▸"):
    """Adds a clean, punchy bullet point item with bold leading title and muted description."""
    p = text_frame.paragraphs[0] if is_first else text_frame.add_paragraph()
    p.space_after = Pt(10)

    # Bullet symbol
    r_icon = p.add_run()
    r_icon.text = f"{icon} "
    r_icon.font.name = FONT_CODE
    r_icon.font.size = Pt(font_size)
    r_icon.font.bold = True
    r_icon.font.color.rgb = accent_color

    # Bold title
    r_title = p.add_run()
    r_title.text = f"{title_bold}: " if title_bold else ""
    r_title.font.name = FONT_HEADING
    r_title.font.size = Pt(font_size)
    r_title.font.bold = True
    r_title.font.color.rgb = TEXT_WHITE

    # Body explanation
    r_desc = p.add_run()
    r_desc.text = description
    r_desc.font.name = FONT_BODY
    r_desc.font.size = Pt(font_size)
    r_desc.font.color.rgb = TEXT_MUTED


def set_speaker_notes(slide, notes_content: str):
    """Sets the slide's speaker notes."""
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_content.strip()


# ==========================================
# SLIDE BUILDERS (8-10 SLIDES CHRONOLOGICAL)
# ==========================================

def build_slide_1_title(prs):
    """Slide 1: Title Slide - Modern hero layout with badges and metrics."""
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Dark background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.color.rgb = BG_COLOR

    # Hero Container Card (Center-Left)
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.9), Inches(11.733), Inches(5.7))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CARD_BORDER
    card.line.width = Pt(1.5)

    # Top Pill Badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.3), Inches(4.3), Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BADGE_BG
    badge.line.color.rgb = CYAN_ACCENT
    badge.line.width = Pt(1)
    btf = badge.text_frame
    btf.margin_top = btf.margin_bottom = btf.margin_left = btf.margin_right = 0
    bp = btf.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    bp.text = "⚡ EXTERNAL ATTACK SURFACE RECONNAISSANCE"
    bp.font.name = FONT_CODE
    bp.font.size = Pt(10)
    bp.font.bold = True
    bp.font.color.rgb = CYAN_ACCENT

    # Main Title & Subtitle Box (Left Column)
    tb = slide.shapes.add_textbox(Inches(1.2), Inches(1.9), Inches(5.8), Inches(2.6))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0

    p_title = tf.paragraphs[0]
    p_title.text = "OmniRecon"
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(48)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE
    p_title.space_after = Pt(4)

    p_tagline = tf.add_paragraph()
    p_tagline.text = "High-Performance Asynchronous Asset Discovery & Vulnerability Footprinting"
    p_tagline.font.name = FONT_HEADING
    p_tagline.font.size = Pt(17)
    p_tagline.font.bold = True
    p_tagline.font.color.rgb = CYAN_ACCENT
    p_tagline.space_after = Pt(10)

    p_desc = tf.add_paragraph()
    p_desc.text = "Bridging Offensive Cybersecurity Reconnaissance with Production-Grade Asynchronous Software Architecture."
    p_desc.font.name = FONT_BODY
    p_desc.font.size = Pt(13)
    p_desc.font.color.rgb = TEXT_MUTED

    # Presenter Information Box (Bottom Left)
    author_box = slide.shapes.add_textbox(Inches(1.2), Inches(4.7), Inches(5.8), Inches(1.5))
    atf = author_box.text_frame
    atf.word_wrap = True
    atf.margin_top = atf.margin_bottom = atf.margin_left = atf.margin_right = 0

    p_by = atf.paragraphs[0]
    p_by.text = "DEVELOPED & PRESENTED BY"
    p_by.font.name = FONT_CODE
    p_by.font.size = Pt(10)
    p_by.font.bold = True
    p_by.font.color.rgb = TEXT_DIM

    p_name = atf.add_paragraph()
    p_name.text = "Deepakraj S"
    p_name.font.name = FONT_HEADING
    p_name.font.size = Pt(22)
    p_name.font.bold = True
    p_name.font.color.rgb = TEXT_WHITE

    p_role = atf.add_paragraph()
    p_role.text = "Software Engineer & Security Researcher\nPython • AsyncIO • Attack Surface Management"
    p_role.font.name = FONT_BODY
    p_role.font.size = Pt(11)
    p_role.font.color.rgb = EMERALD_ACCENT

    # 3 Stat / Capability Chips on the right side (Evenly distributed)
    chips = [
        ("CORE ENGINE ARCHITECTURE", "AsyncIO + aiodns", "Single-threaded non-blocking C-ares event loop handling thousands of concurrent sockets.", CYAN_ACCENT),
        ("DUAL-PARADIGM DISCOVERY", "Passive CT + Active Fuzzing", "Zero-footprint Certificate Transparency paired with wildcard-immune DNS brute forcing.", BLUE_ACCENT),
        ("VULNERABILITY IDENTIFICATION", "Subdomain Takeover Prober", "Automated HTTP/HTTPS status harvesting with dangling cloud bucket fingerprinting.", CORAL_ACCENT),
    ]

    for i, (tag, header, detail, color) in enumerate(chips):
        chip_top = Inches(1.8 + (i * 1.45))
        c_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.3), chip_top, Inches(4.8), Inches(1.25))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG_ALT
        c_shape.line.color.rgb = CARD_BORDER
        c_shape.line.width = Pt(1.2)

        ctb = slide.shapes.add_textbox(Inches(7.5), chip_top + Inches(0.12), Inches(4.4), Inches(1.0))
        ctf = ctb.text_frame
        ctf.word_wrap = True
        ctf.margin_top = ctf.margin_bottom = ctf.margin_left = ctf.margin_right = 0

        cp1 = ctf.paragraphs[0]
        cp1.text = tag
        cp1.font.name = FONT_CODE
        cp1.font.size = Pt(9.5)
        cp1.font.bold = True
        cp1.font.color.rgb = color
        cp1.space_after = Pt(2)

        cp2 = ctf.add_paragraph()
        cp2.text = header
        cp2.font.name = FONT_HEADING
        cp2.font.size = Pt(13)
        cp2.font.bold = True
        cp2.font.color.rgb = TEXT_WHITE
        cp2.space_after = Pt(2)

        cp3 = ctf.add_paragraph()
        cp3.text = detail
        cp3.font.name = FONT_BODY
        cp3.font.size = Pt(10.5)
        cp3.font.color.rgb = TEXT_MUTED

    # Speaker Notes
    set_speaker_notes(slide, """
SPEAKER SCRIPT:
"Good morning, esteemed committee members and recruiters. My name is Deepakraj S, and today I am excited to present OmniRecon: a high-performance, asynchronous reconnaissance toolkit and external attack surface management engine.

In modern cybersecurity, you cannot defend what you don't know exists. OmniRecon addresses a fundamental challenge: how modern enterprises lose track of their exposed infrastructure. 

Equally important to the security problem is how it was engineered. Today, we will examine both the offensive cybersecurity mechanics and the backend systems architecture—specifically why we transitioned away from blocking threads to a high-throughput AsyncIO event loop with non-blocking DNS and HTTP probers. Let's dive in."
""")
    return slide


def build_slide_2_core_problem(prs):
    """Slide 2: The Core Problem - Attack Surface Sprawl & Forgotten Infrastructure."""
    slide = create_base_slide(
        prs,
        category_kicker="// 01. THE CYBERSECURITY PROBLEM",
        slide_title="The Core Problem: External Asset Sprawl & Blind Spots",
        subtitle_takeaway="Organizations cannot defend infrastructure they have forgotten exists."
    )

    # Two prominent side-by-side cards
    # Left Card: Enterprise Reality
    tf_left = add_card(
        slide, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.1),
        title="Why Companies Lose Track", badge_text="THE CAUSE", badge_color=AMBER_ACCENT
    )
    add_bullet_point(tf_left, "Cloud & Microservice Sprawl", "Ephemeral AWS, Azure, and GCP resources spun up without centralized DNS governance.", AMBER_ACCENT, is_first=True)
    add_bullet_point(tf_left, "Shadow IT & Staging Leaks", "Dev teams create temporary testing subdomains (`dev-auth`, `staging-api`) that remain live indefinitely.", AMBER_ACCENT)
    add_bullet_point(tf_left, "M&A Infrastructure Blind Spots", "Acquisitions inherit legacy registrar accounts, unmonitored DNS zones, and untracked endpoints.", AMBER_ACCENT)
    add_bullet_point(tf_left, "Outdated Static Inventory", "Manual asset tracking in spreadsheets goes stale within 24 hours in fast CI/CD environments.", AMBER_ACCENT)

    # Right Card: Security Risks
    tf_right = add_card(
        slide, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.1),
        title="The Technical Threat", badge_text="THE EXPLOITATION", badge_color=CORAL_ACCENT
    )
    add_bullet_point(tf_right, "Dangling DNS Pointers", "CNAME records pointing to deleted cloud buckets or decommissioned third-party services.", CORAL_ACCENT, is_first=True)
    add_bullet_point(tf_right, "Subdomain Takeover", "Adversaries re-register the abandoned cloud resource and serve authenticated phishing on the victim's apex domain.", CORAL_ACCENT)
    add_bullet_point(tf_right, "Unpatched Legacy Services", "Forgotten hosts run outdated software stacks outside enterprise WAF and vulnerability scanner coverage.", CORAL_ACCENT)
    add_bullet_point(tf_right, "Credential Leaks via TXT", "Stray DNS TXT records expose API verification tokens, internal services, and deprecated mail routes.", CORAL_ACCENT)

    # Bottom Callout Card
    tf_bottom = add_card(slide, Inches(0.8), Inches(6.05), Inches(11.7), Inches(0.85), bg_color=CARD_BG_ALT)
    add_bullet_point(
        tf_bottom,
        "THE STRATEGIC REALITY",
        "Adversaries don't break through hardened enterprise firewalls—they discover the forgotten, unpatched subdomain created three years ago by an intern. Defense begins with automated reconnaissance.",
        accent_color=CYAN_ACCENT, is_first=True, icon="⚠️"
    )

    set_speaker_notes(slide, """
SPEAKER SCRIPT:
"Let's look at why external attack surface management is currently one of the top priorities in offensive and defensive security.

Large enterprises rarely have just one domain. In reality, they manage thousands of subdomains across multi-cloud environments. The root problem is asset sprawl and Shadow IT:
Developers spin up an S3 bucket or Heroku app for a quick test, map `staging.company.com` via CNAME, and then delete the cloud resource when finished—without removing the DNS record.

This creates a high-severity vulnerability known as Subdomain Takeover. Because the DNS record is dangling, an attacker can simply register that exact S3 bucket or Heroku app name, instantly hijacking the trusted company subdomain for credential harvesting, session cookie theft, or malware distribution.

Furthermore, static asset inventories fail immediately. Security teams need an automated, high-speed reconnaissance pipeline that continuously discovers assets before attackers exploit them."
""")
    return slide


def build_slide_3_project_overview(prs):
    """Slide 3: Project Overview - High-Level Summary of OmniRecon."""
    slide = create_base_slide(
        prs,
        category_kicker="// 02. SYSTEM OVERVIEW",
        slide_title="Project Overview: What is OmniRecon?",
        subtitle_takeaway="A unified, modular Python reconnaissance engine designed for full external attack surface footprinting."
    )

    # 3 Column Cards
    col_width = Inches(3.7)
    gap = Inches(0.3)

    # Card 1: Modular Pipeline
    tf1 = add_card(
        slide, Inches(0.8), Inches(1.8), col_width, Inches(4.3),
        title="Unified Recon Pipeline", badge_text="PIPELINE", badge_color=CYAN_ACCENT
    )
    add_bullet_point(tf1, "End-to-End Surface Mapping", "Automates the entire reconnaissance lifecycle: from WHOIS ownership to active HTTP probing.", CYAN_ACCENT, is_first=True)
    add_bullet_point(tf1, "Decoupled Architecture", "Independent execution engines: WHOIS, DNS Resolver, Subdomain Enumerator, and Web Prober.", CYAN_ACCENT)
    add_bullet_point(tf1, "Clean Extensibility", "Standardized Python dataclass contracts allow easy integration of custom scanners or security feeds.", CYAN_ACCENT)

    # Card 2: High Concurrency Engine
    tf2 = add_card(
        slide, Inches(0.8 + col_width + gap), Inches(1.8), col_width, Inches(4.3),
        title="AsyncIO Architecture", badge_text="HIGH PERFORMANCE", badge_color=EMERALD_ACCENT
    )
    add_bullet_point(tf2, "Native AsyncIO & aiodns", "Replaces blocking thread pools with non-blocking C-ares DNS queries for order-of-magnitude speedups.", EMERALD_ACCENT, is_first=True)
    add_bullet_point(tf2, "Connection Throttling", "Bounded concurrency via `asyncio.Semaphore` prevents OS file descriptor exhaustion (`EMFILE`).", EMERALD_ACCENT)
    add_bullet_point(tf2, "Wildcard DNS Immunity", "Pre-flight heuristic probing filters out wildcard DNS responses to eliminate false positive noise.", EMERALD_ACCENT)

    # Card 3: Actionable Intelligence
    tf3 = add_card(
        slide, Inches(0.8 + (col_width + gap) * 2), Inches(1.8), col_width, Inches(4.3),
        title="Actionable Intelligence", badge_text="EXPLOIT INSIGHT", badge_color=CORAL_ACCENT
    )
    add_bullet_point(tf3, "Automated Takeover Detection", "Fingerprints dangling cloud providers (AWS S3, GitHub, Heroku, Azure) directly from HTTP bodies.", CORAL_ACCENT, is_first=True)
    add_bullet_point(tf3, "Dual-Port Service Probing", "Concurrently inspects ports 80 & 443 with TLS fallback and automated HTML title harvesting.", CORAL_ACCENT)
    add_bullet_point(tf3, "Enterprise Serialization", "Outputs structured reports in terminal ANSI, machine-readable JSON, and spreadsheet CSV.", CORAL_ACCENT)

    # Bottom Ribbon: Pipeline Stages
    tf_bottom = add_card(slide, Inches(0.8), Inches(6.25), Inches(11.7), Inches(0.68), bg_color=CARD_BG_ALT)
    add_bullet_point(
        tf_bottom,
        "EXECUTION WORKFLOW",
        "[1. WHOIS/RDAP]  ──>  [2. Core DNS Harvesting]  ──>  [3. Passive CT & Active Brute-Force]  ──>  [4. HTTP Service & Takeover Probing]",
        accent_color=BLUE_ACCENT, is_first=True, icon="⚡"
    )

    set_speaker_notes(slide, """
SPEAKER SCRIPT:
"OmniRecon is a modular, high-performance external attack surface footprinting toolkit built entirely in Python.

Rather than stitching together multiple disparate command-line utilities, OmniRecon provides a coherent, four-stage pipeline:
First, it establishes organizational domain ownership via WHOIS and RDAP.
Second, it harvests the apex domain's core DNS records across all major RFC record types.
Third, it enumerates the full attack surface through a dual-paradigm subdomain engine combining passive cryptographic logs with high-concurrency active brute-forcing.
Fourth, it probes live web applications on ports 80 and 443 and inspects them for Subdomain Takeovers.

Architecturally, OmniRecon is built on a clean decoupled design where each module is an independent, testable subsystem producing typed dataclass outputs. Let's examine each phase in detail."
""")
    return slide


def build_slide_4_foundation(prs):
    """Slide 4: Phase 1 & 2 - Foundation: WHOIS Registry & Core DNS Harvesting."""
    slide = create_base_slide(
        prs,
        category_kicker="// 03. RECONNAISSANCE FOUNDATION",
        slide_title="Phase 1 & 2: Domain Ownership & DNS Infrastructure",
        subtitle_takeaway="Establishing legal boundaries, name server authorities, and mail/verification vectors."
    )

    # 2 Column Cards
    col_width = Inches(5.7)

    # Left: Phase 1 - WHOIS / RDAP
    tf_left = add_card(
        slide, Inches(0.8), Inches(1.8), col_width, Inches(4.1),
        title="Phase 1: WHOIS / RDAP Engine", badge_text="THE PROPERTY REGISTRY", badge_color=CYAN_ACCENT
    )
    add_bullet_point(tf_left, "Analogous to Real-Estate Deed", "Identifies who legally owns the target domain, registrar identity, and registration lifespan.", CYAN_ACCENT, is_first=True)
    add_bullet_point(tf_left, "Primary TCP Port 43 Query", "Direct plaintext socket queries to registrar WHOIS servers for low-latency registry extraction.", CYAN_ACCENT)
    add_bullet_point(tf_left, "Automated HTTPS RDAP Fallback", "If Port 43 is firewalled or blocked, automatically fails over to RESTful RDAP over HTTPS Port 443.", CYAN_ACCENT)
    add_bullet_point(tf_left, "Key Metadata Extracted", "Registrar Name & IANA ID, Creation/Expiration dates, Authoritative Nameservers, and EPP Status codes.", CYAN_ACCENT)

    # Right: Phase 2 - Core DNS Harvesting
    tf_right = add_card(
        slide, Inches(6.8), Inches(1.8), col_width, Inches(4.1),
        title="Phase 2: Core DNS Harvester", badge_text="THE MASTER PHONEBOOK", badge_color=EMERALD_ACCENT
    )
    add_bullet_point(tf_right, "Translating Names to Infrastructure", "Queries authoritative records that direct Internet routing and enterprise communication.", EMERALD_ACCENT, is_first=True)
    add_bullet_point(tf_right, "A & AAAA Records", "Maps hostnames to IPv4 and IPv6 endpoints, uncovering origin server infrastructure.", EMERALD_ACCENT)
    add_bullet_point(tf_right, "MX & NS Records", "Identifies primary email gateways (mail routing) and authoritative DNS cluster providers.", EMERALD_ACCENT)
    add_bullet_point(tf_right, "TXT Security Records", "Harvests SPF (`v=spf1`), DMARC (`v=DMARC1`), and third-party SaaS verification tokens (Google, M365).", EMERALD_ACCENT)
    add_bullet_point(tf_right, "SOA Records", "Analyzes Start of Authority serial numbers and zone administrator contact points.", EMERALD_ACCENT)

    # Bottom Architecture Note
    tf_bottom = add_card(slide, Inches(0.8), Inches(6.05), Inches(11.7), Inches(0.85), bg_color=CARD_BG_ALT)
    add_bullet_point(
        tf_bottom,
        "RESILIENT NETWORKING DESIGN",
        "Enterprise firewalls frequently block outbound TCP 43. OmniRecon implements intelligent fault tolerance: native WHOIS attempts timeout within 3s, triggering an immediate non-blocking fallback to ICANN's RESTful RDAP over HTTPS (Port 443).",
        accent_color=BLUE_ACCENT, is_first=True, icon="🛡️"
    )

    set_speaker_notes(slide, """
SPEAKER SCRIPT:
"In Phase 1 and 2, we establish the foundational reconnaissance footprint. I like to explain these using two simple analogies: the 'Property Registry' and the 'Master Phonebook'.

Phase 1 is WHOIS: Think of this as checking the municipal land registry before inspecting a building. It queries TCP port 43 to discover who owns the domain, who the registrar is, when it expires, and which nameservers hold authority. 
A key engineering highlight here is our RDAP fallback. Many corporate and cloud firewalls block port 43. Instead of crashing, OmniRecon detects socket timeouts and automatically fails over to ICANN's modern HTTPS-based RDAP protocol on port 443.

Phase 2 is Core DNS Harvesting: Think of this as the master phonebook translating names into physical addresses. We query dnspython resolvers across A, AAAA, MX, TXT, NS, CNAME, and SOA records.
Crucially, TXT records are an absolute goldmine for security engineers: they reveal email security postures like SPF and DMARC, as well as third-party SaaS verification tokens that prove which cloud platforms the target utilizes."
""")
    return slide


def build_slide_5_subdomain_discovery(prs):
    """Slide 5: Phase 3 - Subdomain Discovery: Passive CT Logs & Active Brute-Forcing."""
    slide = create_base_slide(
        prs,
        category_kicker="// 04. ATTACK SURFACE MAPPING",
        slide_title="Phase 3: Subdomain Discovery (Dual-Paradigm)",
        subtitle_takeaway="Combining zero-footprint cryptographic audit logs with high-concurrency active resolution."
    )

    col_width = Inches(5.7)

    # Left: Passive Certificate Transparency
    tf_left = add_card(
        slide, Inches(0.8), Inches(1.8), col_width, Inches(4.1),
        title="Passive: Certificate Transparency", badge_text="ZERO NETWORK FOOTPRINT", badge_color=CYAN_ACCENT
    )
    add_bullet_point(tf_left, "RFC 6962 Public Audit Logs", "Monitors public, append-only cryptographic logs published by Certificate Authorities (CAs).", CYAN_ACCENT, is_first=True)
    add_bullet_point(tf_left, "Target Undetectable", "Queries third-party aggregators (`crt.sh`) without transmitting a single IP packet to the target network.", CYAN_ACCENT)
    add_bullet_point(tf_left, "Historical Asset Discovery", "Reveals ephemeral, unlinked, or decommissioned subdomains that were ever issued an SSL/TLS certificate.", CYAN_ACCENT)
    add_bullet_point(tf_left, "Instant Seed Generation", "Yields hundreds of validated seed hostnames within milliseconds to prime active scanners.", CYAN_ACCENT)

    # Right: Active High-Concurrency Brute-Forcing
    tf_right = add_card(
        slide, Inches(6.8), Inches(1.8), col_width, Inches(4.1),
        title="Active: AsyncIO Wordlist Engine", badge_text="HIGH-THROUGHPUT C-ARES", badge_color=EMERALD_ACCENT
    )
    add_bullet_point(tf_right, "High-Probability Wordlists", "Brute-forces standard infrastructure names: `api`, `vpn`, `dev`, `admin`, `internal`, `s3`.", EMERALD_ACCENT, is_first=True)
    add_bullet_point(tf_right, "Asynchronous DNS Resolution", "Utilizes `aiodns` and Python's event loop to resolve hundreds of subdomains concurrently.", EMERALD_ACCENT)
    add_bullet_point(tf_right, "Asynchronous Wildcard Immunity", "Pre-probes random pseudo-random subdomains (e.g., `_random991x.target.com`). If wildcard DNS exists, matching IPs are dynamically filtered.", EMERALD_ACCENT)
    add_bullet_point(tf_right, "Resolved Asset Union", "Deduplicates and merges passive CT discoveries with active resolutions into a single verified host list.", EMERALD_ACCENT)

    # Bottom Stat Bar
    tf_bottom = add_card(slide, Inches(0.8), Inches(6.05), Inches(11.7), Inches(0.85), bg_color=CARD_BG_ALT)
    add_bullet_point(
        tf_bottom,
        "WHY DUAL-PARADIGM IS ESSENTIAL",
        "Passive CT logs excel at finding internal naming conventions without touching the target. Active brute-forcing discovers brand-new infrastructure created before certificates are logged. Together, they provide exhaustive coverage with zero wildcard false positives.",
        accent_color=AMBER_ACCENT, is_first=True, icon="🎯"
    )

    set_speaker_notes(slide, """
SPEAKER SCRIPT:
"In Phase 3, we expand the attack surface from a single apex domain to hundreds of potential entry points using a dual-paradigm discovery methodology.

The first paradigm is Passive Certificate Transparency: Under RFC 6962, every SSL certificate issued by CAs like Let's Encrypt or DigiCert must be published to cryptographic, append-only logs. By querying crt.sh, OmniRecon harvests historical and active subdomains with zero network footprint—we do not send a single packet to the target organization.

The second paradigm is Active Brute-Forcing: We take dictionary wordlists and resolve them against DNS at wire speed using our asynchronous `aiodns` engine. 

A critical technical feature here is Wildcard DNS Immunity: If an administrator has a wildcard record like `*.company.com` pointing to a parking page, naive scanners report every word in the wordlist as a valid subdomain, creating thousands of false positives. OmniRecon tests pseudo-random domain strings prior to brute-forcing. If wildcard responses are detected, our engine dynamically blacklists those wildcard IP signatures, ensuring only genuine hostnames are reported."
""")
    return slide


def build_slide_6_http_probing(prs):
    """Slide 6: Phase 4 - HTTP Service Probing & Subdomain Takeover Detection."""
    slide = create_base_slide(
        prs,
        category_kicker="// 05. LIVE SERVICE ANALYSIS & VULNERABILITY AUDITING",
        slide_title="Phase 4: HTTP Service Probing & Subdomain Takeovers",
        subtitle_takeaway="Converting raw DNS resolutions into live application intelligence and flagging high-risk takeovers."
    )

    col_width = Inches(5.7)

    # Left: Web Service Prober
    tf_left = add_card(
        slide, Inches(0.8), Inches(1.8), col_width, Inches(4.1),
        title="HTTP/HTTPS Service Prober", badge_text="PORTS 80 & 443", badge_color=CYAN_ACCENT
    )
    add_bullet_point(tf_left, "Non-Blocking Async Client", "Concurrently probes web endpoints across HTTP (Port 80) and HTTPS (Port 443) via `httpx.AsyncClient`.", CYAN_ACCENT, is_first=True)
    add_bullet_point(tf_left, "Self-Signed TLS Resilience", "Configured with permissive SSL verification (`verify=False`) to ensure internal staging and dev environments respond.", CYAN_ACCENT)
    add_bullet_point(tf_left, "Metadata Harvesting", "Captures HTTP status codes (`200 OK`, `301`, `403`, `500`), web server headers, and redirects.", CYAN_ACCENT)
    add_bullet_point(tf_left, "HTML Title Extraction", "Parses HTML `<title>` tags automatically via regex to give analysts instant contextual visibility into what application is running.", CYAN_ACCENT)

    # Right: Subdomain Takeover Engine
    tf_right = add_card(
        slide, Inches(6.8), Inches(1.8), col_width, Inches(4.1),
        title="Subdomain Takeover Detection", badge_text="HIGH SEVERITY ALERT", badge_color=CORAL_ACCENT
    )
    add_bullet_point(tf_right, "The Mechanics", "Occurs when a DNS CNAME points to an abandoned third-party cloud service that an attacker can claim.", CORAL_ACCENT, is_first=True)
    add_bullet_point(tf_right, "Fingerprint Matching", "Inspects raw response bodies against a curated database of provider error signatures.", CORAL_ACCENT)
    add_bullet_point(tf_right, "Multi-Provider Signatures", "Detects AWS S3 (`NoSuchBucket`), GitHub Pages (`There isn't a GitHub Pages site here`), Heroku (`No such app`), Azure, Shopify, and Fastly.", CORAL_ACCENT)
    add_bullet_point(tf_right, "Instant Alerting", "Flags vulnerable targets with `[!] VULNERABLE` in terminal output and tags them in structured JSON exports.", CORAL_ACCENT)

    # Bottom Visual Flow
    tf_bottom = add_card(slide, Inches(0.8), Inches(6.05), Inches(11.7), Inches(0.85), bg_color=CARD_BG_ALT)
    add_bullet_point(
        tf_bottom,
        "EXPLOITATION CHAIN PREVENTED",
        "Dangling CNAME (`docs.company.com -> company.s3.amazonaws.com`)  ──>  AWS Bucket Deleted  ──>  Attacker Reclaims S3 Bucket  ──>  Complete Domain Hijack & Cookie Theft. OmniRecon catches this instantly.",
        accent_color=CORAL_ACCENT, is_first=True, icon="🚨"
    )

    set_speaker_notes(slide, """
SPEAKER SCRIPT:
"Finding IP addresses is only half the battle. Security analysts and penetration testers need to know: Is there a live web application running, what is it, and is it vulnerable? This is Phase 4.

First, our HTTP Service Prober concurrently inspects every discovered subdomain across Port 80 and Port 443 using `httpx.AsyncClient`. It is resilient to self-signed or expired SSL certificates so internal development hosts aren't dropped. It captures response status codes, web server headers, and automatically parses the HTML `<title>` tag. This allows an analyst to scan a list of 500 subdomains and immediately spot things like 'Grafana Dashboard', 'Jenkins Login', or 'Internal Admin Portal'.

Second, and most importantly, we built automated Subdomain Takeover detection. When a developer points a CNAME to AWS S3, GitHub Pages, or Heroku, and later deletes the cloud instance without cleaning up DNS, the endpoint is vulnerable.
OmniRecon scans response bodies against a curated signature database: if it sees AWS's 'NoSuchBucket' or Heroku's 'No such app', it immediately flags the asset as vulnerable. This turns passive reconnaissance into proactive vulnerability management."
""")
    return slide


def build_slide_7_architecture(prs):
    """Slide 7: Software Engineering Architecture (Crucial Slide) - AsyncIO, Semaphore, Decoupled Design."""
    slide = create_base_slide(
        prs,
        category_kicker="// 06. SYSTEMS ARCHITECTURE (CORE DIFFERENTIATOR)",
        slide_title="Software Architecture: High-Throughput AsyncIO Engine",
        subtitle_takeaway="Architected for maximum concurrency, memory efficiency, and deterministic modularity."
    )

    col_width = Inches(3.7)
    gap = Inches(0.3)

    # 3 Architectural Pillars
    # Pillar 1: Threading vs AsyncIO
    tf1 = add_card(
        slide, Inches(0.8), Inches(1.8), col_width, Inches(4.1),
        title="AsyncIO vs. Threading", badge_text="PERFORMANCE", badge_color=CYAN_ACCENT
    )
    add_bullet_point(tf1, "The Threading Bottleneck", "`ThreadPoolExecutor` allocates OS-level threads. Each incurs a 4–8MB memory stack and heavy kernel context-switching.", CYAN_ACCENT, is_first=True)
    add_bullet_point(tf1, "Single-Thread Event Loop", "`aiodns` interfaces with `c-ares`, an asynchronous DNS library in C, running on Python's native event loop.", CYAN_ACCENT)
    add_bullet_point(tf1, "Massive Resource Efficiency", "Handles thousands of in-flight network queries concurrently with negligible CPU overhead and minimal RAM.", CYAN_ACCENT)

    # Pillar 2: Semaphore & Backpressure
    tf2 = add_card(
        slide, Inches(0.8 + col_width + gap), Inches(1.8), col_width, Inches(4.1),
        title="Connection Throttling", badge_text="RELIABILITY", badge_color=EMERALD_ACCENT
    )
    add_bullet_point(tf2, "The Unbounded Trap", "Unbounded `asyncio.gather()` launches all tasks simultaneously, causing `EMFILE` (Too many open files) crashes.", EMERALD_ACCENT, is_first=True)
    add_bullet_point(tf2, "asyncio.Semaphore Guard", "Strictly controls active in-flight sockets (e.g. `--threads 25`), enforcing non-blocking backpressure.", EMERALD_ACCENT)
    add_bullet_point(tf2, "DNS Rate-Limit Immunity", "Prevents UDP packet loss and DNS resolver drops caused by aggressive bursts on upstream recursive resolvers.", EMERALD_ACCENT)

    # Pillar 3: Decoupled Modularity & Testing
    tf3 = add_card(
        slide, Inches(0.8 + (col_width + gap) * 2), Inches(1.8), col_width, Inches(4.1),
        title="Modular Architecture", badge_text="ENGINEERING RIGOR", badge_color=BLUE_ACCENT
    )
    add_bullet_point(tf3, "Strongly Typed Contracts", "Subsystems return immutable Python `@dataclass` structures (`WhoisResult`, `DnsResult`, `SubdomainResult`).", BLUE_ACCENT, is_first=True)
    add_bullet_point(tf3, "Single Responsibility", "Isolated modules: CLI dispatch, WHOIS lookup, DNS harvester, Subdomain prober, and JSON/CSV exporters.", BLUE_ACCENT)
    add_bullet_point(tf3, "100% Deterministic Testing", "Full test suite using `unittest.mock` and `AsyncMock` ensures zero live network dependencies during CI.", BLUE_ACCENT)

    # Bottom Metric Highlight
    tf_bottom = add_card(slide, Inches(0.8), Inches(6.05), Inches(11.7), Inches(0.85), bg_color=CARD_BG_ALT)
    add_bullet_point(
        tf_bottom,
        "ARCHITECTURAL IMPACT",
        "Replacing synchronous multi-threading with an AsyncIO event loop and Semaphore throttling eliminated thread starvation, reduced memory consumption by ~85%, and enabled seamless scanning of 10,000+ domain wordlists.",
        accent_color=EMERALD_ACCENT, is_first=True, icon="🚀"
    )

    set_speaker_notes(slide, """
SPEAKER SCRIPT:
"Now I want to focus on what I consider the most crucial slide of this presentation: the software engineering architecture behind OmniRecon. For recruiters and engineers, this is where the toolkit truly shines.

When building reconnaissance tools, developers often start with standard multi-threading using `ThreadPoolExecutor`. However, OS threads carry heavy memory overhead—4 to 8 megabytes per thread stack—and incur constant kernel context switching. Scaling to hundreds of threads leads to thread starvation and high latency.

We solved this by rebuilding OmniRecon around Python's native `asyncio` event loop paired with `aiodns`. `aiodns` interfaces directly with `c-ares`, an asynchronous C-based DNS library. This allows thousands of network queries to run non-blocking on a single operating system thread.

However, high concurrency introduces another classic systems challenge: socket exhaustion. If you fire 10,000 DNS queries simultaneously via an unbounded `asyncio.gather()`, your OS crashes with `EMFILE`—Too Many Open Files—and DNS servers drop UDP packets.
We resolved this with `asyncio.Semaphore`. The semaphore acts as a gatekeeper, dynamically throttling concurrent connections to a controlled threshold while maintaining maximum saturation.

Finally, the entire codebase follows clean code principles: modular engines decoupled via strongly typed dataclasses and verified by a 100% mocked offline test suite utilizing `AsyncMock`."
""")
    return slide


def build_slide_8_future_enhancements(prs):
    """Slide 8: Future Enhancements - CI/CD Pipelines, Dockerization, AXFR Zone Transfers."""
    slide = create_base_slide(
        prs,
        category_kicker="// 07. FUTURE ROADMAP",
        slide_title="Future Enhancements & Engineering Roadmap",
        subtitle_takeaway="Extending OmniRecon into continuous automated security monitoring and DevOps pipelines."
    )

    col_width = Inches(3.7)
    gap = Inches(0.3)

    # 3 Enhancement Cards
    # Enhancement 1: CI/CD & Containerization
    tf1 = add_card(
        slide, Inches(0.8), Inches(1.8), col_width, Inches(4.1),
        title="CI/CD & Containerization", badge_text="DEVSECOPS PIPELINE", badge_color=CYAN_ACCENT
    )
    add_bullet_point(tf1, "Docker Multi-Stage Build", "Package OmniRecon into a minimal, lightweight OCI container image for reproducible deployments.", CYAN_ACCENT, is_first=True)
    add_bullet_point(tf1, "GitHub Actions Integration", "Automated scheduled workflow that scans enterprise domains weekly and alerts teams to changes.", CYAN_ACCENT)
    add_bullet_point(tf1, "Attack Surface Drift Detection", "Diffs current scan results against historical JSON runs to alert on newly exposed subdomains in real-time.", CYAN_ACCENT)

    # Enhancement 2: DNS Zone Transfer (AXFR)
    tf2 = add_card(
        slide, Inches(0.8 + col_width + gap), Inches(1.8), col_width, Inches(4.1),
        title="DNS Zone Transfer (AXFR)", badge_text="OFFENSIVE AUDITING", badge_color=CORAL_ACCENT
    )
    add_bullet_point(tf2, "Authoritative NS Auditing", "Automatically query discovered nameservers for unauthenticated DNS Zone Transfer (`AXFR`) misconfigurations.", CORAL_ACCENT, is_first=True)
    add_bullet_point(tf2, "Full Zone Extraction", "If a misconfigured nameserver permits AXFR, download and parse the entire internal DNS zone in seconds.", CORAL_ACCENT)
    add_bullet_point(tf2, "Asynchronous Implementation", "Integrate non-blocking AXFR queries directly into the DNS harvesting phase.", CORAL_ACCENT)

    # Enhancement 3: Dynamic Permutation Engine
    tf3 = add_card(
        slide, Inches(0.8 + (col_width + gap) * 2), Inches(1.8), col_width, Inches(4.1),
        title="Permutation Fuzzing", badge_text="INTELLIGENT ENUMERATION", badge_color=EMERALD_ACCENT
    )
    add_bullet_point(tf3, "Contextual Mutations", "Extract naming patterns (`api-prod`, `dev-vpn`) and dynamically generate targeted alterations (`api-dev`, `api-uat`).", EMERALD_ACCENT, is_first=True)
    add_bullet_point(tf3, "Recursive Async Resolution", "Feed generated permutations back into the `aiodns` engine to uncover unlinked development clusters.", EMERALD_ACCENT)
    add_bullet_point(tf3, "Executive HTML/PDF Reports", "Jinja2-powered standalone executive summaries with risk scoring and visual network topology graphs.", EMERALD_ACCENT)

    # Bottom Callout Banner
    tf_bottom = add_card(slide, Inches(0.8), Inches(6.05), Inches(11.7), Inches(0.85), bg_color=CARD_BG_ALT)
    add_bullet_point(
        tf_bottom,
        "VISION",
        "Transforming OmniRecon from a point-in-time CLI scanner into a continuous, cloud-native Attack Surface Management (ASM) daemon for modern DevSecOps environments.",
        accent_color=BLUE_ACCENT, is_first=True, icon="🔮"
    )

    set_speaker_notes(slide, """
SPEAKER SCRIPT:
"While OmniRecon is already a robust standalone toolkit, our engineering roadmap includes three high-value extensions:

First, CI/CD and Containerization: We are creating a minimal Docker container and a GitHub Actions workflow. This enables continuous attack surface monitoring—by diffing JSON outputs between scheduled runs, security teams can detect 'attack surface drift' the moment a developer exposes a new staging domain.

Second, DNS Zone Transfer (AXFR) Testing: An authoritative nameserver should never permit unrestricted AXFR requests. We will automate asynchronous AXFR queries against all discovered NS records. If a misconfigured nameserver leaks its entire zone file, OmniRecon will ingest all internal records instantly.

Third, Dynamic Permutation Fuzzing: Instead of relying purely on static wordlists, our permutation engine will identify contextual patterns—like `api-prod`—and automatically generate permutations like `api-staging`, `api-dev`, and `api-test` for targeted discovery.

This roadmap evolves OmniRecon from a manual utility into a continuous DevSecOps reconnaissance daemon."
""")
    return slide


def build_slide_9_conclusion(prs):
    """Slide 9: Q&A / Conclusion - Summary, Tech Stack, Contact & Q&A Invitation."""
    slide = create_base_slide(
        prs,
        category_kicker="// 08. CONCLUSION & DISCUSSION",
        slide_title="Conclusion & Technical Q&A",
        subtitle_takeaway="A high-performance bridge between offensive security discovery and modern async software engineering."
    )

    col_width = Inches(5.7)

    # Left: Key Engineering Achievements
    tf_left = add_card(
        slide, Inches(0.8), Inches(1.8), col_width, Inches(4.1),
        title="Key Technical Takeaways", badge_text="PROJECT SUMMARY", badge_color=CYAN_ACCENT
    )
    add_bullet_point(tf_left, "Full Attack Surface Coverage", "Unifies WHOIS/RDAP, core DNS records, passive CT logs, active brute-force, and HTTP probing.", CYAN_ACCENT, is_first=True)
    add_bullet_point(tf_left, "High-Throughput Concurrency", "AsyncIO event loop + aiodns + httpx delivers enterprise-grade scalability with bounded memory usage.", CYAN_ACCENT)
    add_bullet_point(tf_left, "Actionable Vulnerability Impact", "Directly identifies critical Subdomain Takeovers across major cloud infrastructure providers.", CYAN_ACCENT)
    add_bullet_point(tf_left, "Software Engineering Best Practices", "Clean decoupled architecture, typed dataclasses, resilient fallbacks, and 100% offline mocked unit tests.", CYAN_ACCENT)

    # Right: Technical Stack & Q&A Card
    tf_right = add_card(
        slide, Inches(6.8), Inches(1.8), col_width, Inches(4.1),
        title="Open for Questions", badge_text="TECH STACK & REPO", badge_color=EMERALD_ACCENT
    )
    add_bullet_point(tf_right, "Core Language", "Python 3.9+ (Modern type hinting, AsyncIO event loop)", EMERALD_ACCENT, is_first=True)
    add_bullet_point(tf_right, "Asynchronous Networking", "`aiodns` (C-ares DNS resolver), `httpx` (Async HTTP/2 client)", EMERALD_ACCENT)
    add_bullet_point(tf_right, "DNS & Network Libraries", "`dnspython`, `python-whois`, ICANN RESTful RDAP fallback", EMERALD_ACCENT)
    add_bullet_point(tf_right, "Testing & Code Quality", "`unittest`, `unittest.mock.AsyncMock`, `flake8`, `black`", EMERALD_ACCENT)
    add_bullet_point(tf_right, "Presenter", "Deepakraj S  |  GitHub: OmniRecon Project", EMERALD_ACCENT)

    # Bottom Ribbon: Invitation
    tf_bottom = add_card(slide, Inches(0.8), Inches(6.05), Inches(11.7), Inches(0.85), bg_color=CARD_BG_ALT)
    add_bullet_point(
        tf_bottom,
        "THANK YOU",
        "I welcome any questions regarding the AsyncIO architecture, concurrency throttling trade-offs, takeover detection algorithms, or future roadmap items.",
        accent_color=CYAN_ACCENT, is_first=True, icon="💬"
    )

    set_speaker_notes(slide, """
SPEAKER SCRIPT:
"To conclude: OmniRecon demonstrates how core computer networking, offensive cybersecurity reconnaissance, and high-performance software engineering converge.

By replacing blocking thread architectures with an AsyncIO event loop and semaphore throttling, we achieved massive throughput gains while maintaining socket safety and eliminating wildcard false positives. By coupling passive CT logs with automated Subdomain Takeover detection, we transformed raw DNS resolution into actionable security intelligence.

Thank you very much for your time and attention. I am now delighted to open the floor to your questions regarding the architecture, technical trade-offs, or future extensions."
""")
    return slide


def generate_presentation(output_filepath="OmniRecon_Presentation.pptx"):
    """Main presentation generator routine."""
    print("[*] Initializing PowerPoint presentation...")
    prs = Presentation()

    # Configure 16:9 widescreen dimensions (13.333" x 7.5")
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    print("[*] Building Slide 1: Title Slide...")
    build_slide_1_title(prs)

    print("[*] Building Slide 2: The Core Problem...")
    build_slide_2_core_problem(prs)

    print("[*] Building Slide 3: Project Overview...")
    build_slide_3_project_overview(prs)

    print("[*] Building Slide 4: Phase 1 & 2 - Foundation...")
    build_slide_4_foundation(prs)

    print("[*] Building Slide 5: Phase 3 - Subdomain Discovery...")
    build_slide_5_subdomain_discovery(prs)

    print("[*] Building Slide 6: Phase 4 - HTTP Service Probing...")
    build_slide_6_http_probing(prs)

    print("[*] Building Slide 7: Software Engineering Architecture...")
    build_slide_7_architecture(prs)

    print("[*] Building Slide 8: Future Enhancements...")
    build_slide_8_future_enhancements(prs)

    print("[*] Building Slide 9: Conclusion & Q&A...")
    build_slide_9_conclusion(prs)

    # Save presentation
    prs.save(output_filepath)
    print(f"[+] Presentation successfully generated: {os.path.abspath(output_filepath)}")
    return output_filepath


if __name__ == "__main__":
    out_file = "OmniRecon_Presentation.pptx"
    if len(sys.argv) > 1:
        out_file = sys.argv[1]
    generate_presentation(out_file)
